"""Shared paths, table schemas, logging, and file helpers."""

from __future__ import annotations

import hashlib
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping, Sequence

import pandas as pd
import yaml
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


RAW_COLUMNS = [
    "Paper_ID（论文编号）",
    "Title（标题）",
    "Authors（作者）",
    "Abstract（摘要）",
    "Venue（会议/期刊）",
    "Year（年份）",
    "Track（论文轨道）",
    "Official_URL（官方论文页面）",
    "PDF_URL（官方PDF链接）",
    "BibTeX（BibTeX信息）",
    "Crawl_Status（抓取状态）",
    "Crawl_Error（抓取错误）",
    "Notes（备注）",
]

SCREENING_COLUMNS = [
    "Paper_ID（论文编号）",
    "Title（标题）",
    "Authors（作者）",
    "Abstract（摘要）",
    "Venue（会议/期刊）",
    "Year（年份）",
    "Track（论文轨道）",
    "Rule_Score（规则分数）",
    "Positive_Hits（命中的阳性规则）",
    "Negative_Hits（命中的负向规则）",
    "Teacher_Evidence（Teacher证据）",
    "Student_Evidence（Student证据）",
    "Transfer_Evidence（知识迁移证据）",
    "Large_Model_Evidence（大模型相关证据）",
    "Efficient_AI_Evidence（高效AI相关证据）",
    "Decision（筛选决定）",
    "Decision_Reason（筛选理由）",
    "Need_Manual_Check（是否需要人工复核）",
    "Official_URL（官方论文页面）",
    "PDF_URL（官方PDF链接）",
    "Download_Status（下载状态）",
    "Local_PDF_Path（本地PDF路径）",
    "Rule_Version（规则版本）",
    "Notes（备注）",
]

MANIFEST_COLUMNS = [
    "Paper_ID（论文编号）",
    "Title（标题）",
    "Decision（筛选决定）",
    "PDF_URL（官方PDF链接）",
    "Local_PDF_Path（本地PDF路径）",
    "Download_Status（下载状态）",
    "Download_Error（下载错误）",
    "File_Size（文件大小）",
    "Download_Time（下载时间）",
]

DECISION_KEEP = "RULE_KEEP（规则保留）"
DECISION_MAYBE = "RULE_MAYBE（规则待复核）"
DECISION_AMBIGUOUS = "RULE_AMBIGUOUS（规则冲突）"
DECISION_DROP = "SAFE_DROP（安全排除）"


@dataclass(frozen=True)
class ProjectPaths:
    """Resolved project and venue-year output paths."""

    root: Path
    venue: str
    year: int

    @property
    def output_root(self) -> Path:
        return self.root / "output" / f"{self.venue.upper()}_{self.year}"

    @property
    def raw(self) -> Path:
        return self.output_root / "raw"

    @property
    def cache(self) -> Path:
        return self.raw / "cache"

    @property
    def screening(self) -> Path:
        return self.output_root / "screening"

    @property
    def pdfs(self) -> Path:
        return self.output_root / "PDFs"

    @property
    def manifests(self) -> Path:
        return self.output_root / "manifests"

    @property
    def reports(self) -> Path:
        return self.output_root / "reports"

    @property
    def logs(self) -> Path:
        return self.output_root / "logs"

    @property
    def stem(self) -> str:
        return f"{self.venue.upper()}{self.year}"

    @property
    def raw_csv(self) -> Path:
        return self.raw / f"{self.stem}_All_Papers.csv"

    @property
    def raw_xlsx(self) -> Path:
        return self.raw / f"{self.stem}_All_Papers.xlsx"

    @property
    def screening_csv(self) -> Path:
        return self.screening / f"{self.stem}_KD_Screening.csv"

    @property
    def screening_xlsx(self) -> Path:
        return self.screening / f"{self.stem}_KD_Screening.xlsx"

    def ensure(self) -> None:
        for directory in (
            self.raw,
            self.cache / "lists",
            self.cache / "details",
            self.screening,
            self.pdfs / "RULE_KEEP",
            self.pdfs / "RULE_MAYBE",
            self.manifests,
            self.reports,
            self.logs,
        ):
            directory.mkdir(parents=True, exist_ok=True)


def load_yaml(path: Path) -> dict:
    """Load a UTF-8 YAML mapping."""
    with path.open("r", encoding="utf-8") as handle:
        data = yaml.safe_load(handle)
    if not isinstance(data, dict):
        raise ValueError(f"YAML root must be a mapping: {path}")
    return data


def setup_logger(name: str, log_file: Path) -> logging.Logger:
    """Create a console and UTF-8 file logger without duplicate handlers."""
    log_file.parent.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger(name)
    logger.setLevel(logging.INFO)
    logger.propagate = False
    if logger.handlers:
        return logger
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    file_handler = logging.FileHandler(log_file, encoding="utf-8")
    file_handler.setFormatter(formatter)
    console_handler = logging.StreamHandler()
    console_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    logger.addHandler(console_handler)
    return logger


def stable_paper_id(venue: str, year: int, track: str, url: str) -> str:
    """Build a stable, compact ID from immutable source identity."""
    track_code = {"Main Conference": "MAIN", "Findings": "FIND", "Workshop": "WS"}.get(
        track, re.sub(r"\W+", "", track.upper())[:8]
    )
    digest = hashlib.sha1(url.encode("utf-8")).hexdigest()[:10].upper()
    return f"{venue.upper()}{year}_{track_code}_{digest}"


def write_csv(rows: Iterable[Mapping[str, object]] | pd.DataFrame, path: Path, columns: Sequence[str]) -> None:
    """Write an Excel-friendly UTF-8 BOM CSV with a stable schema."""
    path.parent.mkdir(parents=True, exist_ok=True)
    frame = rows.copy() if isinstance(rows, pd.DataFrame) else pd.DataFrame(list(rows))
    frame = frame.reindex(columns=columns, fill_value="")
    frame.to_csv(path, index=False, encoding="utf-8-sig", lineterminator="\n")


def write_xlsx(rows: Iterable[Mapping[str, object]] | pd.DataFrame, path: Path, columns: Sequence[str]) -> None:
    """Write a formatted workbook with frozen header, filter, widths, and wrapped abstracts."""
    path.parent.mkdir(parents=True, exist_ok=True)
    frame = rows.copy() if isinstance(rows, pd.DataFrame) else pd.DataFrame(list(rows))
    frame = frame.reindex(columns=columns, fill_value="")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Papers")
        sheet = writer.book["Papers"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for idx, column in enumerate(columns, start=1):
            if column == "Abstract（摘要）":
                width = 80
            elif column in {"BibTeX（BibTeX信息）", "Decision_Reason（筛选理由）"}:
                width = 55
            elif "URL" in column or "Path" in column:
                width = 48
            else:
                sample_lengths = [len(str(column))]
                sample_lengths.extend(len(str(value)) for value in frame[column].head(200).fillna(""))
                width = min(max(sample_lengths) + 2, 35)
            sheet.column_dimensions[get_column_letter(idx)].width = max(width, 12)
            if column in {"Abstract（摘要）", "BibTeX（BibTeX信息）", "Decision_Reason（筛选理由）"}:
                for cell in sheet[get_column_letter(idx)][1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)


def clean_cell(value: object) -> str:
    """Convert pandas/CSV values to a predictable string."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value)

